#!/usr/bin/env python3
"""
LOM Consolidator — Transform multi-sheet LOM (List of Materials) Excel files
into Mohaimen-style single-sheet format with Persian section headers.

Transforms data center Bill-of-Materials spreadsheets from the common
multi-sheet Iranian procurement format into a consolidated single-sheet
layout suitable for printing / review.

STRUCTURAL TRANSFORMATION:
  1. CONSOLIDATE → All sheets merged into one sheet
  2. SECTIONS → Each sheet → merged A:F header (22pt bold "بخش ...")
     with data rows below (18pt regular, center-center, wrap=ON)
  3. HEADERS → Single row (12pt bold, theme color, no explicit fill)
  4. COLUMNS → A-F (ردیف, تجهیزات, توضیحات, مقدار, واحد, برند)
               + G-N empty template with dark blue fill (FF0070C0) on row 1
  5. BORDERS → Thin on all sides for every cell
  6. ROW GROUPING → All data rows outline_level=1 per section
  7. PAGE SETUP → Landscape A4, print area = A:F

Usage:
  python lom_consolidate.py <input.xlsx> [options]

Options:
  --output <path>     Output path (default: ~/Desktop/<stem>-Mohaimen.xlsx)
  --sheet-name <s>    Target sheet name (default: cleaned filename)
  --order <order>     Comma-separated section order by sheet-name substring
                      e.g. "MECH,CIVIL,ACS,POWER,FAS,FES,CCTV,PASSIVE"

Examples:
  python lom_consolidate.py "LOM-IT01.xlsx"
  python lom_consolidate.py "LOM-IT01.xlsx" --sheet-name "IT-01" \
    --order "MECH,CIVIL,ACS,POWER,FAS,FES,CCTV,PASSIVE"
"""

import sys, os, re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# ── Style constants ──────────────────────────────────────────────────────────

THIN = Side(style='thin')
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITLE_FONT    = Font(name='Calibri', size=20, bold=True)
HEADER_FONT   = Font(name='Calibri', size=12, bold=True)
SECTION_FONT  = Font(name='Calibri', size=22, bold=True)
DATA_FONT     = Font(name='Calibri', size=18, bold=False)
CENTER_WRAP   = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Dark blue fill for template column headers (G1:N1)
TEMPLATE_FILL = PatternFill(start_color='FF0070C0', end_color='FF0070C0', fill_type='solid')

COL_HEADERS   = ['ردیف', 'تجهیزات', 'توضیحات', 'مقدار', 'واحد', 'برند']

COL_WIDTHS = {
    'A': 8.6640625, 'B': 112.0,       'C': 41.6640625, 'D': 12.5546875,
    'E': 14.33203125, 'F': 22.33203125,
    'G': 9.109375, 'H': 13.0, 'I': 13.0, 'J': 13.0,
    'K': 13.0, 'L': 13.0, 'M': 13.0, 'N': 13.0,
}

# ── Persian section-name mapping ─────────────────────────────────────────────

SECTION_NAMES = {
    'civil': 'بخش عمران', 'mech': 'بخش مکانیک', 'mechanical': 'بخش مکانیک',
    'power': 'بخش POWER',
    'acs': 'بخش اکسس کنترل',
    'cctv': 'بخش CCTV',
    'passive': 'بخش PASSIVE',
    'fas': 'بخش FAS', 'firefighting': 'بخش FAS',
    'fes': 'بخش FES',
}

def derive_section_name(sheet_name):
    """Convert 'CIVIL-LOM' → 'بخش عمران'."""
    stem = re.sub(r'[-_\s]*(LOM|LOS)\s*$', '', sheet_name.strip(), flags=re.I).strip()
    key = stem.lower().replace('-', '').replace(' ', '').replace('_', '')
    return SECTION_NAMES.get(key, f'بخش {stem}')


# ── Helpers ──────────────────────────────────────────────────────────────────

def _cell(ws, row, col, value, font, alignment, border=THIN_BORDER):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font; c.alignment = alignment; c.border = border
    return c

def _write_header(ws, row):
    for i, h in enumerate(COL_HEADERS, 1):
        _cell(ws, row, i, h, HEADER_FONT, CENTER_WRAP)
    for c in range(7, 15):
        _cell(ws, row, c, None, HEADER_FONT, CENTER_WRAP)

def _write_section(ws, row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    _cell(ws, row, 1, text, SECTION_FONT, CENTER_WRAP)
    for c in range(2, 15):
        _cell(ws, row, c, None, SECTION_FONT, CENTER_WRAP)

def _write_data(ws, row, values):
    padded = list(values) + [None] * (6 - len(values))
    for i, v in enumerate(padded, 1):
        _cell(ws, row, i, v, DATA_FONT, CENTER_WRAP)
    for c in range(7, 15):
        _cell(ws, row, c, None, DATA_FONT, CENTER_WRAP)

def _read_data(src, sname):
    """Return list of data rows (cols A-F) skipping title + header."""
    ws = src[sname]
    return [list(r[:6]) for r in ws.iter_rows(min_row=3, values_only=True)
            if any(v is not None for v in r)]


# ── Main transform ───────────────────────────────────────────────────────────

def consolidate(input_path, output_path, sheet_name=None, section_order=None):
    """Read multi-sheet LOM → write single-sheet Mohaimen-style file."""

    src = load_workbook(input_path)
    dst = Workbook()
    ws = dst.active

    # Sheet name
    stem = os.path.splitext(os.path.basename(input_path))[0]
    ws.title = (sheet_name or re.sub(r'[-\s]+', '_', stem))[:31]

    # Column widths
    for letter, w in COL_WIDTHS.items():
        ws.column_dimensions[letter].width = w

    # Section ordering
    all_sheets = src.sheetnames
    if section_order:
        ordered = []
        for item in section_order:
            item = item.strip()
            if item.isdigit():
                idx = int(item)
                if 0 <= idx < len(all_sheets):
                    ordered.append(all_sheets[idx])
            else:
                ordered.extend(s for s in all_sheets if item.lower() in s.lower())
        seen = set(ordered)
        ordered.extend(s for s in all_sheets if s not in seen)
    else:
        ordered = list(all_sheets)

    row = 1

    # 1. Title row (merged A:F, all 14 cols bordered)
    title = src[ordered[0]].cell(row=1, column=1).value or 'لیست تجهیزات'
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    _cell(ws, 1, 1, title, TITLE_FONT, CENTER_WRAP)
    for c in range(2, 15):
        _cell(ws, 1, c, None, TITLE_FONT, CENTER_WRAP)
    # Dark blue fill on G1:N1 (template column headers)
    for c in range(7, 15):
        ws.cell(row=1, column=c).fill = TEMPLATE_FILL
    row += 1

    # 2. Column headers
    _write_header(ws, row)
    row += 1

    # 3. Sections with row grouping
    for sname in ordered:
        data = _read_data(src, sname)
        if not data:
            continue
        _write_section(ws, row, derive_section_name(sname))
        section_start = row + 1
        row += 1
        for d in data:
            _write_data(ws, row, d)
            row += 1
        # Group data rows under this section (outline_level=1)
        for r in range(section_start, row):
            ws.row_dimensions[r].outline_level = 1

    # 4. Page setup: landscape A4, print area A:F
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9  # A4
    ws.print_area = f"'{ws.title}'!$A$1:$F${row - 1}"

    dst.save(output_path)
    src.close()
    dst.close()
    print(f"✅ → {output_path}  ({row-1} rows, 14 cols)")
    return output_path


# ── CLI entry point ──────────────────────────────────────────────────────────

if __name__ == '__main__':
    if len(sys.argv) < 2 or sys.argv[1] in ('-h', '--help'):
        print(__doc__)
        sys.exit(0)

    input_path = sys.argv[1]
    if not os.path.exists(input_path):
        print(f"❌ File not found: {input_path}")
        sys.exit(1)

    output_path = sheet_name = None
    section_order = None
    args = sys.argv[2:]
    while args:
        arg = args.pop(0)
        if arg == '--output' and args:        output_path = args.pop(0)
        elif arg == '--sheet-name' and args:  sheet_name = args.pop(0)
        elif arg == '--order' and args:       section_order = [x.strip() for x in args.pop(0).split(',')]

    if not output_path:
        stem = os.path.splitext(os.path.basename(input_path))[0]
        output_path = os.path.join(os.path.expanduser('~'), 'Desktop', f'{stem}-Mohaimen.xlsx')

    consolidate(input_path, output_path, sheet_name, section_order)
